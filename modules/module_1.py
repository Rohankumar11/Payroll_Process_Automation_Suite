"""
Module 1 - Payroll Report Generation Engine.

This module wraps the original working automation script (module_!.py)
into callable functions WITHOUT changing the business logic.

Business workflow:
    1. Read Higher Pension master file (Approved + Rejected sheets).
    2. Map UAN to Member ID.
    3. Read raw EPFO wage files (tab-delimited .xls exports).
    4. Clean wage data and format wage month.
    5. Round EPF values using floor(x + 0.5).
    6. Generate employee-wise payroll workbooks with formulas,
       borders, bold headers, footer block and auto column widths.
    7. Write success / error / missing member logs.

This module must never import Streamlit. UI progress is reported
through an optional callback function.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side

# Optional callback used by the UI layer: callback(stage_text, fraction)
ProgressCallback = Optional[Callable[[str, float], None]]

APPROVED_SHEET = "Approved Application"
REJECTED_SHEET = "Rejected Application"

WAGES_MONTH_FORMATS = ("%m/%Y", "%b-%y", "%d-%m-%Y", "%d/%m/%Y")


# ----------------------------------------------------------------------
# Internal helpers
# ----------------------------------------------------------------------
def _report(callback: ProgressCallback, stage: str, fraction: float) -> None:
    """Safely forward a progress update to the UI layer."""
    if callback is not None:
        callback(stage, fraction)


def _write_lines(path: Path, lines: List[str]) -> None:
    """Write log lines exactly like the original script (one per line)."""
    with open(path, "w", encoding="utf-8") as handle:
        for line in lines:
            handle.write(line + "\n")


def format_wages_month(value: Any) -> str:
    """
    Normalize a wage month value to 'MMM-YY' format.

    Preserves the original multi-format parsing behavior.
    Falls back to the raw value when no format matches.
    """
    value = str(value).strip()
    for fmt in WAGES_MONTH_FORMATS:
        try:
            return pd.to_datetime(value, format=fmt).strftime("%b-%y")
        except Exception:
            pass
    return value


def _load_master_lookup(master_file: Path) -> Dict[str, str]:
    """
    Build the UAN -> Member ID lookup dictionary.

    Preserves original cleaning logic:
        - Keep only UAN / Member ID columns
        - Concat Approved + Rejected sheets
        - Drop empty UANs
        - Remove '.0' artifacts and whitespace
        - Drop duplicate UANs
    """
    approved_df = pd.read_excel(master_file, sheet_name=APPROVED_SHEET)
    rejected_df = pd.read_excel(master_file, sheet_name=REJECTED_SHEET)

    approved_clean = approved_df[["UAN", "Member ID"]]
    rejected_clean = rejected_df[["UAN", "Member ID"]]

    master_df = pd.concat([approved_clean, rejected_clean], ignore_index=True)
    master_df = master_df[["UAN", "Member ID"]]
    master_df = master_df.dropna(subset=["UAN"])
    master_df["UAN"] = (
        master_df["UAN"]
        .astype(str)
        .str.replace(".0", "", regex=False)
        .str.strip()
    )
    master_df["Member ID"] = master_df["Member ID"].astype(str).str.strip()
    master_df = master_df.drop_duplicates(subset=["UAN"])

    return dict(zip(master_df["UAN"], master_df["Member ID"]))


def _build_employee_frame(
    employee_df: pd.DataFrame,
    uan_to_member: Dict[str, str],
) -> pd.DataFrame:
    """
    Build the per-employee output dataframe.

    Column order is business-critical:
        A Member Name | B Member ID | C Wages Month | D EPF Wages |
        E Arrear | F EPS Diversion | G Revised Salary | H NCP Days

    The G-column Excel formula (=D+E) depends on this order.
    """
    output_df = pd.DataFrame()
    output_df["Member Name"] = employee_df["Member Name"]
    output_df["Member ID"] = (
        employee_df["UAN of the member/applicant"].map(uan_to_member)
    )
    output_df["Wages Month"] = (
        employee_df["Wages Month"].apply(format_wages_month)
    )
    output_df["EPF Wages"] = (
        pd.to_numeric(employee_df["EPF Wages"], errors="coerce")
        .fillna(0)
        .apply(lambda x: int(np.floor(x + 0.5)))
    )
    output_df["Arrear"] = 0
    output_df["EPS Diversion"] = (
        pd.to_numeric(
            employee_df["EPS Contribution (Pension A/C)"],
            errors="coerce",
        )
        .fillna(0)
        .apply(lambda x: int(np.floor(x + 0.5)))
    )
    output_df["Revised Salary"] = ""
    output_df["NCP Days"] = 0
    return output_df


def _write_employee_workbook(
    output_df: pd.DataFrame,
    output_path: Path,
) -> None:
    """
    Write one employee workbook with the original formatting:
        - G column formula =D+E
        - Thin borders on all cells
        - Bold headers
        - Summary footer block at max_row + 3
        - Auto column width (length + 5)
    """
    wb = Workbook()
    ws = wb.active

    # Headers
    ws.append(list(output_df.columns))

    # Data rows + G column formula
    start_row = 2
    for idx, row in enumerate(output_df.values.tolist(), start=start_row):
        ws.append(row)
        ws[f"G{idx}"] = f"=D{idx}+E{idx}"

    # Borders
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border

    # Bold headers
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # Footer block
    footer_row = ws.max_row + 3
    ws[f"B{footer_row}"] = "DATA UPLOADED IN EPFO"
    ws[f"C{footer_row}"] = "FINAL DATA BY TSL"
    ws[f"D{footer_row}"] = "DIFFERENCE"
    ws[f"E{footer_row}"] = "REMARKS"

    # Auto column width
    for column_cells in ws.columns:
        length = max(
            len(str(cell.value)) if cell.value else 0
            for cell in column_cells
        )
        ws.column_dimensions[column_cells[0].column_letter].width = length + 5

    wb.save(output_path)


# ----------------------------------------------------------------------
# Public API
# ----------------------------------------------------------------------
def generate_reports(
    master_file: Path,
    wage_files: Sequence[Path],
    output_folder: Path,
    log_folder: Path,
    progress_callback: ProgressCallback = None,
) -> Dict[str, Any]:
    """
    Module 1 entry point.

    Args:
        master_file:
            Higher Pension master workbook containing
            'Approved Application' and 'Rejected Application' sheets.
        wage_files:
            Raw EPFO wage files (tab-delimited .xls exports).
        output_folder:
            Destination folder for employee workbooks.
        log_folder:
            Destination folder for success/error/missing logs.
        progress_callback:
            Optional UI callback: callback(stage_text, fraction).

    Returns:
        Execution summary dictionary.

    Raises:
        ValueError / KeyError:
            For missing sheets or missing master columns.
            The UI layer converts these into user-friendly messages.
    """
    started_at = datetime.now()

    output_folder = Path(output_folder)
    log_folder = Path(log_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    log_folder.mkdir(parents=True, exist_ok=True)

    success_logs: List[str] = []
    error_logs: List[str] = []
    missing_member_logs: List[str] = []

    # ------------------------------------------------------------------
    # MASTER FILE
    # ------------------------------------------------------------------
    _report(progress_callback, "Reading Master File", 0.05)
    uan_to_member = _load_master_lookup(Path(master_file))
    _report(progress_callback, "Mapping Member IDs", 0.15)

    # ------------------------------------------------------------------
    # WAGE FILES (original filter: only *.xls)
    # ------------------------------------------------------------------
    wage_paths = sorted(
        Path(p) for p in wage_files if str(p).endswith(".xls")
    )
    total_files = max(len(wage_paths), 1)

    reports_generated = 0
    employees_processed = set()

    for file_index, file_path in enumerate(wage_paths):
        file_name = file_path.name
        try:
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            _report(
                progress_callback,
                f"Generating Employee Reports — {file_name}",
                0.2 + 0.7 * (file_index / total_files),
            )

            # Original read behavior: EPFO .xls exports are
            # tab-delimited text files, NOT binary Excel files.
            df = pd.read_csv(
                file_path,
                sep="\t",
                dtype={"UAN of the member/applicant": str},
            )

            df["UAN of the member/applicant"] = (
                df["UAN of the member/applicant"]
                .astype(str)
                .str.replace(".0", "", regex=False)
                .str.strip()
            )

            unique_uans = df["UAN of the member/applicant"].unique()
            total_uans = max(len(unique_uans), 1)

            for uan_index, uan in enumerate(unique_uans):
                employee_df = df[
                    df["UAN of the member/applicant"] == uan
                ]

                _report(
                    progress_callback,
                    f"Generating Employee Reports — UAN {uan}",
                    0.2 + 0.7 * ((file_index + uan_index / total_uans) / total_files),
                )

                output_df = _build_employee_frame(
                    employee_df, uan_to_member
                )

                # Missing Member ID logging (report still generated)
                if output_df["Member ID"].isna().any():
                    missing_member_logs.append(
                        f"[{timestamp}] "
                        f"{file_name} --> "
                        f"Member ID NOT FOUND for UAN {uan}"
                    )

                _report(
                    progress_callback,
                    f"Formatting Reports — UAN {uan}",
                    0.2 + 0.7 * ((file_index + (uan_index + 1) / total_uans) / total_files),
                )

                output_path = output_folder / f"{uan}.xlsx"

                # Duplicate UAN handling: log, then overwrite
                if output_path.exists():
                    error_logs.append(
                        f"[{timestamp}] "
                        f"DUPLICATE UAN FOUND -> {uan} "
                        f"in file {file_name}"
                    )

                _write_employee_workbook(output_df, output_path)

                reports_generated += 1
                employees_processed.add(uan)

            success_logs.append(f"[{timestamp}] {file_name} --> SUCCESS")

        except Exception as exc:  # per-file resilience (original behavior)
            error_logs.append(
                f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] "
                f"{file_name} --> ERROR: {exc}"
            )

    # ------------------------------------------------------------------
    # LOGS (original file names)
    # ------------------------------------------------------------------
    _report(progress_callback, "Writing Logs", 0.92)
    _write_lines(log_folder / "success_log.txt", success_logs)
    _write_lines(log_folder / "error_log.txt", error_logs)
    _write_lines(log_folder / "missing_memberid.txt", missing_member_logs)

    duration = (datetime.now() - started_at).total_seconds()

    return {
        "module": "Module 1 — Payroll Report Generation",
        "status": "success",
        "employees_processed": len(employees_processed),
        "reports_generated": reports_generated,
        "missing_members": len(missing_member_logs),
        "errors": len(error_logs),
        "files_successful": len(success_logs),
        "duration_seconds": round(duration, 2),
        "output_folder": str(output_folder),
        "log_folder": str(log_folder),
        "success_logs": success_logs,
        "error_logs": error_logs,
        "missing_member_logs": missing_member_logs,
    }