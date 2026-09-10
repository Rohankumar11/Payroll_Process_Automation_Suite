"""
Utility module - OPR Basic Pay Filler.

Fills the "Basic" column of an OPR wages workbook from a basic-pay
matrix workbook, and produces a missing-data report.

Design rules:
    - Basic matrix: ALL sheets read (names never hard-coded).
      Header row detected via a "P No" marker; month columns detected
      dynamically from date headers (new months append automatically).
    - P.No values are normalized everywhere (strip spaces / \\u00A0,
      commas, quotes -> int) because exact-match lookups fail on the
      raw text values.
    - #N/A / blank matrix cells are skipped.
    - Wages months parsed from real dates OR "Mmm-yy" text.
    - Months before the matrix range (Apr-2005) are legitimately
      missing: left untouched and reported.
    - No new sheets created; no other columns modified;
      formatting preserved (openpyxl in-memory edit).

Manual acceptance check (developer console):
    from modules import basic_pay_filler as bpf
    m = bpf.build_basic_lookup(open("Basic_OPR.xlsx", "rb").read())
    assert m["lookup"][(121796, 2005, 4)] == 21800
    assert m["lookup"][(120324, 2006, 8)] == 32400
    assert m["lookup"][(120324, 2011, 4)] == 105300
    assert all(k[0] != 121418 for k in m["lookup"])
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

MonthKey = Tuple[int, int]

EXCEL_EPOCH = pd.Timestamp("1899-12-30")
INVALID_TOKENS = {"", "#N/A", "N/A", "NA", "-", "#REF!", "#VALUE!", "NULL"}

# Fallback column positions when headers are absent:
# P.No(A) Month(B) UAN(C) Name(D) Basic(J)
FALLBACK_COLUMNS = {"pno": 1, "month": 2, "uan": 3, "name": 4, "basic": 10}

HEADER_ALIASES = {
    "pno": {"pno"},
    "month": {"wagesmonth", "wagemonth", "month"},
    "uan": {"uan"},
    "name": {"membername", "name"},
    "basic": {"basic", "basicpay"},
}


# ----------------------------------------------------------------------
# Normalization helpers
# ----------------------------------------------------------------------
def normalize_header(value: Any) -> str:
    """Lowercase a header and remove spaces/dots/underscores/nbsp."""
    if value is None:
        return ""
    text = str(value).lower()
    text = text.replace("\u00A0", "").replace(".", "").replace("_", "")
    return re.sub(r"\s+", "", text)


def normalize_pno(value: Any) -> Optional[int]:
    """
    Normalize a P.No stored as int, float or dirty text
    (leading/trailing spaces, non-breaking spaces, commas).
    Returns None for blanks / #N/A / unparsable values.
    """
    if value is None:
        return None
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        if np.isnan(value):
            return None
        return int(value) if float(value).is_integer() else None
    text = str(value).replace("\u00A0", "").replace(",", "")
    text = text.strip().strip('"').strip("'").strip()
    if not text or text.upper() in INVALID_TOKENS or text.startswith("#"):
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


def _to_number(value: Any) -> Optional[float]:
    """Convert a matrix pay cell to int/float; None for blank/#N/A."""
    if value is None:
        return None
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        if np.isnan(value):
            return None
        return int(value) if float(value).is_integer() else float(value)
    if isinstance(value, str):
        text = value.replace("\u00A0", "").replace(",", "").strip()
        if not text or text.upper() in INVALID_TOKENS or text.startswith("#"):
            return None
        try:
            number = float(text)
        except ValueError:
            return None
        return int(number) if number.is_integer() else number
    return None


def _header_to_month(value: Any) -> Optional[MonthKey]:
    """Parse a basic-matrix header cell into (year, month)."""
    if isinstance(value, pd.Timestamp):
        return (value.year, value.month) if pd.notna(value) else None
    if isinstance(value, datetime):
        return (value.year, value.month)
    if isinstance(value, (int, float, np.integer, np.floating)):
        number = float(value)
        if np.isnan(number) or not 15000 < number < 80000:
            return None
        stamp = EXCEL_EPOCH + pd.Timedelta(days=number)
        return (stamp.year, stamp.month)
    if isinstance(value, str):
        stamp = pd.to_datetime(value.strip(), errors="coerce", dayfirst=True)
        return (stamp.year, stamp.month) if pd.notna(stamp) else None
    return None


def parse_wages_month(value: Any) -> Tuple[Optional[MonthKey], str]:
    """
    Parse a wages 'Wages Month' cell (real date or 'Mmm-yy' text).

    Returns ((year, month) or None, display label).
    """
    if value is None:
        return None, ""
    if isinstance(value, (pd.Timestamp, datetime)):
        stamp = pd.Timestamp(value)
        if pd.isna(stamp):
            return None, str(value)
        return (stamp.year, stamp.month), stamp.strftime("%b-%y")
    text = str(value).strip()
    for fmt in ("%b-%y", "%B-%y", "%b-%Y", "%m/%Y", "%m-%y"):
        try:
            stamp = pd.to_datetime(text, format=fmt)
            return (stamp.year, stamp.month), stamp.strftime("%b-%y")
        except Exception:
            continue
    stamp = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if pd.isna(stamp):
        return None, text
    return (stamp.year, stamp.month), stamp.strftime("%b-%y")


# ----------------------------------------------------------------------
# Basic matrix lookup
# ----------------------------------------------------------------------
def build_basic_lookup(basic_bytes: bytes) -> Dict[str, Any]:
    """
    Read ALL sheets of the basic-pay matrix (header=None) and build:

        lookup : {(p_no, year, month): basic_pay}
        pnos   : set of all P.No values present in the matrix
        months : sorted list of detected (year, month) columns
    """
    sheets = pd.read_excel(
        io.BytesIO(basic_bytes), sheet_name=None, header=None
    )

    lookup: Dict[Tuple[int, int, int], float] = {}
    pnos: set = set()
    months: set = set()

    for _sheet_name, frame in sheets.items():
        if frame.empty:
            continue
        array = frame.to_numpy()
        rows, cols = array.shape

        # ---- locate header row + P.No column ----
        header_row = None
        pno_col = None
        for r in range(min(10, rows)):
            for c in range(min(6, cols)):
                if normalize_header(array[r, c]) == "pno":
                    header_row, pno_col = r, c
                    break
            if header_row is not None:
                break
        if header_row is None:
            continue

        # ---- detect month columns dynamically ----
        month_cols: Dict[int, MonthKey] = {}
        for c in range(pno_col + 1, cols):
            ym = _header_to_month(array[header_row, c])
            if ym is not None:
                month_cols[c] = ym
                months.add(ym)

        # ---- data rows ----
        for r in range(header_row + 1, rows):
            pno = normalize_pno(array[r, pno_col])
            if pno is None:
                continue
            pnos.add(pno)
            for c, ym in month_cols.items():
                number = _to_number(array[r, c])
                if number is None:
                    continue
                key = (pno, ym[0], ym[1])
                if key not in lookup:
                    lookup[key] = number

    return {"lookup": lookup, "pnos": pnos, "months": sorted(months)}


# ----------------------------------------------------------------------
# Wages workbook column detection
# ----------------------------------------------------------------------
def _detect_columns(ws: Worksheet) -> Tuple[Dict[str, int], int]:
    """
    Detect header row and column indexes by header text,
    falling back to fixed positions (A,B,C,D,J) when absent.
    """
    for r in range(1, min(10, ws.max_row + 1)):
        found: Dict[str, int] = {}
        for c in range(1, min(20, ws.max_column + 1)):
            header = normalize_header(ws.cell(r, c).value)
            for key, aliases in HEADER_ALIASES.items():
                if header in aliases and key not in found:
                    found[key] = c
        if "pno" in found:
            for key, col in FALLBACK_COLUMNS.items():
                found.setdefault(key, col)
            return found, r
    return dict(FALLBACK_COLUMNS), 1


# ----------------------------------------------------------------------
# Report builder
# ----------------------------------------------------------------------
def _build_report(
    total: int,
    filled: int,
    missed: int,
    member_summary: List[Dict[str, Any]],
    row_details: List[Dict[str, Any]],
    range_label: str,
) -> str:
    """Build the Not_Found_Data.txt content."""
    bar = "=" * 78
    thin = "-" * 78
    lines = [
        bar,
        "NOT FOUND DATA REPORT",
        f"Generated          : {datetime.now():%Y-%m-%d %H:%M:%S}",
        f"Basic matrix range : {range_label}",
        bar,
        f"Total rows processed : {total}",
        f"Rows filled          : {filled}",
        f"Rows not found       : {missed}",
        f"Members with missing : {len(member_summary)}",
        "",
        thin,
        "MEMBER SUMMARY",
        thin,
        "P.No | Member Name | Total Rows | Filled | Not Found | Status",
    ]
    for member in member_summary:
        lines.append(
            f"{member['P.No']} | {member['Member Name']} | "
            f"{member['Total Rows']} | {member['Filled']} | "
            f"{member['Not Found']} | {member['Status']}"
        )
    lines += ["", thin, "ROW DETAILS", thin,
              "Sheet | Row | P.No | Member Name | Wages Month"]
    for detail in row_details:
        lines.append(
            f"{detail['Sheet']} | {detail['Row']} | {detail['P.No']} | "
            f"{detail['Member Name']} | {detail['Wages Month']}"
        )
    lines.append("")
    return "\n".join(lines)


# ----------------------------------------------------------------------
# Public API
# ----------------------------------------------------------------------
def fill_basic_pay(basic_bytes: bytes, wages_bytes: bytes) -> Dict[str, Any]:
    """
    Fill the Basic column of the wages workbook from the basic matrix.

    Args:
        basic_bytes: Raw bytes of Basic_OPR.xlsx.
        wages_bytes: Raw bytes of the OPR wages workbook.

    Returns:
        Dictionary with keys:
            workbook  : bytes of the updated wages workbook
            report    : Not_Found_Data.txt content
            filled    : rows filled
            missed    : rows not found
            total     : rows processed
            member_summary / row_details : structured data for the UI
            range_label : detected matrix range e.g. 'Apr-05..Jun-26'
    """
    basic = build_basic_lookup(basic_bytes)
    lookup = basic["lookup"]
    pnos = basic["pnos"]
    months = basic["months"]

    if months:
        low = datetime(months[0][0], months[0][1], 1)
        high = datetime(months[-1][0], months[-1][1], 1)
        range_label = f"{low:%b-%y}..{high:%b-%y}"
    else:
        range_label = "n/a"

    wb = load_workbook(io.BytesIO(wages_bytes))

    total = filled = missed = 0
    members: Dict[Any, Dict[str, Any]] = {}
    row_details: List[Dict[str, Any]] = []

    for ws in wb.worksheets:
        mapping, header_row = _detect_columns(ws)

        for row_num in range(header_row + 1, ws.max_row + 1):
            pno_raw = ws.cell(row_num, mapping["pno"]).value
            month_raw = ws.cell(row_num, mapping["month"]).value

            # skip completely empty rows
            if pno_raw is None and month_raw is None:
                continue

            total += 1
            pno = normalize_pno(pno_raw)
            ym, month_label = parse_wages_month(month_raw)
            name = ws.cell(row_num, mapping["name"]).value

            member_key = pno if pno is not None else f"RAW:{pno_raw}"
            stats = members.setdefault(
                member_key,
                {"pno": pno, "name": name or "", "total": 0,
                 "filled": 0, "missed": 0},
            )
            if name and not stats["name"]:
                stats["name"] = name
            stats["total"] += 1

            value = (
                lookup.get((pno, ym[0], ym[1]))
                if (pno is not None and ym is not None)
                else None
            )

            if value is not None:
                ws.cell(row_num, mapping["basic"]).value = value
                filled += 1
                stats["filled"] += 1
            else:
                missed += 1
                stats["missed"] += 1
                row_details.append(
                    {
                        "Sheet": ws.title,
                        "Row": row_num,
                        "P.No": pno if pno is not None else str(pno_raw),
                        "Member Name": name or "",
                        "Wages Month": month_label,
                    }
                )

    buffer = io.BytesIO()
    wb.save(buffer)

    # ---------------- member summary ----------------
    member_summary: List[Dict[str, Any]] = []
    ordered = sorted(
        members.values(),
        key=lambda s: (s["pno"] is None, s["pno"] or 0),
    )
    for stats in ordered:
        if stats["missed"] == 0:
            continue
        if stats["pno"] is None:
            status = "P.No missing/invalid in wages row"
        elif stats["pno"] not in pnos:
            status = "P.No NOT present in Basic file"
        else:
            status = (
                f"P.No IS in Basic file ({range_label}); "
                "missing months outside range or blank/#N/A"
            )
        member_summary.append(
            {
                "P.No": stats["pno"],
                "Member Name": stats["name"],
                "Total Rows": stats["total"],
                "Filled": stats["filled"],
                "Not Found": stats["missed"],
                "Status": status,
            }
        )

    report = _build_report(
        total, filled, missed, member_summary, row_details, range_label
    )

    return {
        "workbook": buffer.getvalue(),
        "report": report,
        "filled": filled,
        "missed": missed,
        "total": total,
        "member_summary": member_summary,
        "row_details": row_details,
        "range_label": range_label,
    }