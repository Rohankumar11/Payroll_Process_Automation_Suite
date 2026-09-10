"""
Module 2 - Stage 2: Payroll Validation Engine.

Wraps the original working script (module_2_2.py) into a callable
function WITHOUT changing business logic.

Workflow per processed workbook:
    1. Read salary values from 'Processed_Data'.
    2. Recalculate Revised Salary / Arrear:
         - if Arrear > 0  -> Revised = EPF + Arrear
         - else          -> Arrear = Revised - EPF
    3. Remove duplicate summary blocks (rewrite one clean block).
    4. Recalculate Final TSL Total and Difference.
    5. Save the validated workbook.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from openpyxl import load_workbook

ProgressCallback = Optional[Callable[[str, float], None]]

SUMMARY_MARKER = "DATA UPLOADED IN EPFO"


def _log(log_folder: Optional[Path], message: str, is_error: bool = False) -> None:
    """Append validation log lines when a log folder is provided."""
    if log_folder is None:
        return
    log_folder = Path(log_folder)
    log_folder.mkdir(parents=True, exist_ok=True)
    log_file = log_folder / "validation_log.txt"
    tag = "[ERROR]" if is_error else "[INFO]"
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(log_file, "a", encoding="utf-8") as handle:
        handle.write(f"{tag} {timestamp} - {message}\n")


def validate_reports(
    processed_folder: Path,
    log_folder: Optional[Path] = None,
    progress_callback: ProgressCallback = None,
) -> Dict[str, Any]:
    """
    Module 2 Stage 2 entry point.

    Validates every .xlsx workbook inside processed_folder in place
    and returns a per-file reconciliation summary.
    """
    started_at = datetime.now()
    processed_folder = Path(processed_folder)

    file_results: List[Dict[str, Any]] = []
    validated_count = 0
    corrected_count = 0
    error_count = 0

    xlsx_files = sorted(processed_folder.glob("*.xlsx"))
    total = max(len(xlsx_files), 1)

    for position, file_path in enumerate(xlsx_files):
        file_name = file_path.name
        uan = file_name.replace("PROCESSED_", "").replace(".xlsx", "")

        if progress_callback:
            progress_callback(
                f"Validating {file_name}",
                0.05 + 0.9 * (position / total),
            )

        try:
            _log(log_folder, f"Processing : {file_name}")
            wb = load_workbook(file_path)
            ws = wb["Processed_Data"]

            # ---------------- column letters ----------------
            headers = {}
            for cell in ws[1]:
                if cell.value is not None:
                    headers[str(cell.value).strip()] = cell.column_letter
            epf_col = headers["EPF Wages"]
            arrear_col = headers["Arrear"]
            revised_col = headers["Revised Salary"]

            # ---------------- last valid summary ----------------
            summary_row = None
            for row_num in range(ws.max_row, 1, -1):
                if ws[f"B{row_num}"].value == SUMMARY_MARKER:
                    summary_row = row_num
                    break

            epfo_total = 0.0
            remarks = ""
            if summary_row:
                try:
                    epfo_total = float(ws[f"B{summary_row + 1}"].value)
                except Exception:
                    epfo_total = 0
                remarks = ws[f"E{summary_row + 1}"].value
                if remarks is None:
                    remarks = ""

            end_row = summary_row - 1 if summary_row else ws.max_row

            # ---------------- recalculation ----------------
            final_tsl = 0.0
            corrected = False
            for row_num in range(2, end_row + 1):
                month_value = ws[f"C{row_num}"].value
                if month_value in [None, ""]:
                    continue
                try:
                    epf = float(ws[f"{epf_col}{row_num}"].value or 0)
                except Exception:
                    epf = 0
                try:
                    arrear = float(ws[f"{arrear_col}{row_num}"].value or 0)
                except Exception:
                    arrear = 0
                try:
                    revised = float(ws[f"{revised_col}{row_num}"].value or 0)
                except Exception:
                    revised = 0

                if arrear > 0:
                    revised = epf + arrear
                    ws[f"{revised_col}{row_num}"] = revised
                    corrected = True
                else:
                    arrear = revised - epf
                    ws[f"{arrear_col}{row_num}"] = arrear
                    corrected = True

                final_tsl += revised

            # ---------------- remove duplicate summaries ----------------
            first_summary = None
            for row_num in range(1, ws.max_row + 1):
                if ws[f"B{row_num}"].value == SUMMARY_MARKER:
                    first_summary = row_num
                    break
            if first_summary:
                ws.delete_rows(first_summary, ws.max_row - first_summary + 1)

            # ---------------- last data row ----------------
            last_data_row = 1
            for row_num in range(ws.max_row, 1, -1):
                if ws[f"C{row_num}"].value not in [None, ""]:
                    last_data_row = row_num
                    break

            difference = final_tsl - epfo_total
            start_row = last_data_row + 3

            ws[f"B{start_row}"] = "DATA UPLOADED IN EPFO"
            ws[f"C{start_row}"] = "FINAL DATA BY TSL"
            ws[f"D{start_row}"] = "DIFFERENCE"
            ws[f"E{start_row}"] = "REMARKS"
            ws[f"B{start_row + 1}"] = epfo_total
            ws[f"C{start_row + 1}"] = final_tsl
            ws[f"D{start_row + 1}"] = difference
            ws[f"E{start_row + 1}"] = remarks

            wb.save(file_path)

            validated_count += 1
            if corrected:
                corrected_count += 1

            file_results.append(
                {
                    "UAN": uan,
                    "Status": "Validated",
                    "Corrected": corrected,
                    "EPFO Total": round(epfo_total, 2),
                    "Final TSL Total": round(final_tsl, 2),
                    "Difference": round(difference, 2),
                    "Remarks": remarks,
                }
            )
            _log(log_folder, f"Updated : {file_name}")

        except Exception as exc:
            error_count += 1
            file_results.append(
                {"UAN": uan, "Status": "Error", "Corrected": False, "Reason": str(exc)}
            )
            _log(log_folder, f"Error : {file_name} : {exc}", is_error=True)

    if progress_callback:
        progress_callback("Validation Completed", 1.0)

    duration = (datetime.now() - started_at).total_seconds()

    return {
        "stage": "Stage 2 — Validation",
        "files": file_results,
        "validated": validated_count,
        "corrected": corrected_count,
        "errors": error_count,
        "duration_seconds": round(duration, 2),
    }